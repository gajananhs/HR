#!/usr/bin/env python3
"""
Import CAPL's source spreadsheets into the Quickair HR PWA.

  python3 scripts/import-data.py --master Emp_Name_With_ID.xlsx \
      --salary salary_statement.xlsx --skills skill_matrix_*.xlsx \
      --payroll-out ../private/payroll-2026-07.json --report-out ../private/import-report.json

Writes (inside the app):
  js/data-employees.js   employee master (name + code only — the source has nothing else)
  js/data-skills.js      every skill matrix sheet, employees matched to codes where unambiguous
Writes (OUTSIDE the app — never commit / deploy):
  payroll JSON           wages register rows, for the encrypted on-device import in the Pay screen
  report JSON            every validation flag raised during import

Nothing is invented: blank cells stay null, and anything that cannot be matched is flagged.
Re-run this whenever the source spreadsheets change.
"""
import argparse, datetime as dt, json, re, sys
from pathlib import Path
import openpyxl
from openpyxl.utils import column_index_from_string as ci

ROOT = Path(__file__).resolve().parent.parent
flags = []  # {area, severity, item, detail}


def flag(area, severity, item, detail):
    flags.append({'area': area, 'severity': severity, 'item': item, 'detail': detail})


def clean(v):
    if isinstance(v, str):
        v = re.sub(r'\s+', ' ', v).strip()
        return v or None
    return v


# ------------------------------------------------------------------ employee master
def read_master(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    hdr = None
    for r in range(1, 10):
        vals = [str(ws.cell(r, c).value or '').lower() for c in range(1, 6)]
        if any('name' in v for v in vals) and any('code' in v or 'id' in v for v in vals):
            hdr = r
            name_c = next(c for c in range(1, 6) if 'name' in str(ws.cell(r, c).value or '').lower())
            code_c = next(c for c in range(1, 6) if re.search(r'code|id', str(ws.cell(r, c).value or '').lower()))
            break
    if not hdr:
        sys.exit('Employee master: header row not found')
    emps, seen_codes, seen_names = [], {}, {}
    for r in range(hdr + 1, ws.max_row + 1):
        name, code = clean(ws.cell(r, name_c).value), clean(ws.cell(r, code_c).value)
        if not name and not code:
            continue
        if not name or not code:
            flag('Employee master', 'high', f'Row {r}', f'Incomplete record (name={name!r}, code={code!r}) — not imported')
            continue
        code = str(code).upper()
        if not re.fullmatch(r'CAPL_\d{4}', code):
            flag('Employee master', 'medium', code, f'Unexpected code format in row {r}')
        if code in seen_codes:
            flag('Employee master', 'high', code, f'Duplicate code (rows {seen_codes[code]} and {r}) — second row skipped')
            continue
        key = name.lower()
        if key in seen_names:
            flag('Employee master', 'medium', name, f'Same name appears twice ({seen_names[key]} and {code})')
        seen_codes[code] = r
        seen_names[key] = code
        emps.append({'id': code, 'name': name})
    return emps


# ------------------------------------------------------------------ salary register
SAL_COLS = {
    'slNo': 'A', 'name': 'B', 'doj': 'C', 'uan': 'D', 'wageDays': 'E', 'payDays': 'F', 'otHrs': 'G', 'otRate': 'H',
    'ctc': 'I', 'fixedWages': 'J',
    'fixed.basicVda': 'K', 'fixed.hra': 'L', 'fixed.ca': 'M', 'fixed.ma': 'N', 'fixed.sa': 'O', 'fixed.bonus': 'P',
    'pfWages': 'Q',
    'earned.basicVda': 'R', 'earned.hra': 'S', 'earned.ca': 'T', 'earned.ma': 'U', 'earned.sa': 'V', 'earned.bonus': 'W',
    'earned.ot': 'X', 'earned.attnBonus': 'Y', 'earned.otherEarnings': 'Z', 'earned.arrearsIncentive': 'AA',
    'gross': 'AB', 'esiWages': 'AC',
    'ded.esi': 'AD', 'ded.epf': 'AE', 'ded.pt': 'AF', 'ded.alreadyPaid': 'AG', 'ded.attDed': 'AH', 'ded.salAdv': 'AI', 'ded.tds': 'AJ',
    'totalDed': 'AK', 'net': 'AL',
}


def num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    try:
        return round(float(str(v).replace(',', '')), 2)
    except ValueError:
        return None


def date_iso(v, where):
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    m = re.fullmatch(r'(\d{1,2})[./-](\d{1,2})[./-](\d{4})', str(v).strip())
    if m:
        d, mo, y = map(int, m.groups())
        return dt.date(y, mo, d).isoformat()
    flag('Salary register', 'medium', where, f'Unreadable DOJ {v!r}')
    return None


def read_salary(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    title = ' '.join(str(ws.cell(r, 1).value or '') for r in range(1, 4))
    m = re.search(r'MONTH OF\s+([A-Z]+)\s+(\d{4})', title.upper())
    if not m:
        sys.exit('Salary register: month not found in title')
    month = dt.datetime.strptime(f'{m.group(1)[:3]} {m.group(2)}', '%b %Y').strftime('%Y-%m')
    company = clean(ws.cell(1, 1).value)
    address = clean(ws.cell(2, 1).value)

    rows = []
    for r in range(6, ws.max_row + 1):
        label = clean(ws.cell(r, ci('B')).value)
        if label and label.lower() == 'total':
            totals_row = r
            break
        if ws.cell(r, 1).value is None:
            continue
        rec = {'row': r, 'earned': {}, 'fixed': {}, 'ded': {}}
        for key, col in SAL_COLS.items():
            raw = ws.cell(r, ci(col)).value
            if key == 'doj':
                val = date_iso(raw, f'Row {r}')
            elif key in ('name', 'uan'):
                val = clean(raw)
                val = str(val) if val is not None else None
            elif key == 'slNo':
                val = int(raw)
            else:
                val = num(raw)
            if '.' in key:
                grp, sub = key.split('.')
                rec[grp][sub] = val
            else:
                rec[key] = val
        rows.append(rec)
    else:
        totals_row = None

    # ---- validation
    where = lambda x: f"Sl {x['slNo']} (row {x['row']})"
    named = [x for x in rows if x['name']]
    if not named:
        flag('Salary register', 'high', 'All rows',
             f"Employee's Name column is blank in all {len(rows)} rows (UAN on {sum(1 for x in rows if x['uan'])} only) — "
             'salary cannot be linked to employee codes from the file; map each Sl No in the Pay screen')
    sl = [x['slNo'] for x in rows]
    dups = sorted({s for s in sl if sl.count(s) > 1})
    if dups:
        flag('Salary register', 'high', 'Sl No', f'Duplicate Sl No: {dups}')
    missing = sorted(set(range(1, max(sl) + 1)) - set(sl))
    if missing:
        flag('Salary register', 'medium', 'Sl No', f'Sequence gap — Sl No {missing} not present; order is {sl[-3:]} at the end')
    for x in rows:
        if not x['doj']:
            flag('Salary register', 'low', where(x), 'DOJ blank')
        e_sum = round(sum(v or 0 for v in x['earned'].values()), 2)
        if x['gross'] is not None and abs(e_sum - x['gross']) > 1:
            flag('Salary register', 'medium', where(x), f"Earned components sum {e_sum} ≠ Gross {x['gross']}")
        d_sum = round(sum(v or 0 for v in x['ded'].values()), 2)
        if x['totalDed'] is not None and abs(d_sum - x['totalDed']) > 1:
            flag('Salary register', 'medium', where(x), f"Deductions sum {d_sum} ≠ Total Ded {x['totalDed']}")
        if None not in (x['gross'], x['totalDed'], x['net']) and abs(x['gross'] - x['totalDed'] - x['net']) > 1:
            flag('Salary register', 'medium', where(x), f"Gross − Total Ded ≠ Net ({x['gross']} − {x['totalDed']} vs {x['net']})")
        f_sum = round(sum(v or 0 for v in x['fixed'].values()), 2)
        if x['fixedWages'] is not None and abs(f_sum - x['fixedWages']) > 1:
            flag('Salary register', 'low', where(x), f"Fixed components sum {f_sum} ≠ Fixed Wages {x['fixedWages']}")
        for k, v in x['earned'].items():
            if v is not None and v != int(v):
                flag('Salary register', 'low', where(x), f'Fractional amount in {k}: {v}')
        if x['payDays'] is not None and x['wageDays'] and x['payDays'] > x['wageDays']:
            flag('Salary register', 'medium', where(x), 'Pay days exceed wage days')
    if totals_row:
        for key in ('gross', 'totalDed', 'net'):
            col = SAL_COLS[key]
            sheet_total = num(ws.cell(totals_row, ci(col)).value)
            calc = round(sum(x[key] or 0 for x in rows), 2)
            if sheet_total is not None and abs(sheet_total - calc) > 1:
                flag('Salary register', 'medium', f'Total {key}', f'Sheet total {sheet_total} ≠ sum of rows {calc}')
    extra = [r for r in range(ci('AM'), ws.max_column + 1)
             if any(ws.cell(rr, r).value not in (None, '') for rr in range(6, ws.max_row + 1))]
    if extra or (totals_row and ws.max_row > totals_row):
        flag('Salary register', 'low', 'Side notes',
             'Working notes in columns AM–AR and below the Total row (e.g. "kumta", "direcotrs", "jayalakshmi", "sathish", "pintu") '
             'are not employee rows — ignored, not imported')
    return {'month': month, 'company': company, 'address': address, 'source': Path(path).name,
            'columns': 'Fixed = monthly structure; Earned = pro-rated for pay days, as in the register',
            'rows': rows}


# ------------------------------------------------------------------ skill matrices
LEGEND = {1: 'Trainee', 2: 'Perform based on instructions', 3: 'Can guide others', 4: 'Expert'}
STOP_NAMES = {'description', 'trainee', 'perform based on instructions', 'can guide others', 'expert'}


def merged_top(ws, r, c):
    for m in ws.merged_cells.ranges:
        if m.min_row <= r <= m.max_row and m.min_col <= c <= m.max_col:
            return ws.cell(m.min_row, m.min_col).value
    return ws.cell(r, c).value


def level(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip().upper()
        if s == 'NA':
            return 'NA'
        if s.isdigit():
            v = int(s)
        else:
            return None
    return int(v) if float(v).is_integer() else float(v)


DEPT_LABEL = {'[MACHINESHOP]': 'Machine Shop', 'PRODUCTION': None}


def read_skill_book(path):
    out = []
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        title = ws.title.strip()
        hdr = None
        for r in range(1, 15):
            for c in range(1, 8):
                v = ws.cell(r, c).value
                if isinstance(v, str) and 'employee name' in v.lower():
                    hdr = (r, c)
                    break
            if hdr:
                break
        if not hdr:
            # Sheet3 in every skill workbook is an unrelated ID/Names/Category/Designation list (IDs 4–75,
            # e.g. "General Manager - ASIA") that matches no CAPL employee — skipped and flagged once.
            if ws.max_row > 1 and str(ws.cell(1, 3).value or '').upper() == 'NAMES':
                flag('Skill matrix', 'low', f'{Path(path).name} › {title}',
                     'Unrelated staff list (IDs 4–75, e.g. "KALIRAJ GURUSAMY", "General Manager - ASIA") — no CAPL code/name match; not imported')
            continue
        hr, nc = hdr
        cols = {}
        for c in range(1, ws.max_column + 1):
            v = str(ws.cell(hr, c).value or '').lower()
            if 'employee id' in v:
                cols['id'] = c
            elif 'job title' in v or 'designation' in v:
                cols['title'] = c
        dept = None
        for r in range(1, hr):
            v = ws.cell(r, 1).value
            if isinstance(v, str) and 'department' in v.lower():
                dept = v.split(':', 1)[-1].strip()
        raw_dept = (dept or '').upper()
        file_label = Path(path).stem.split('-', 1)[-1]
        if raw_dept in DEPT_LABEL and DEPT_LABEL[raw_dept]:
            dept = DEPT_LABEL[raw_dept]
        elif raw_dept == 'PRODUCTION' or not dept:
            # Two workbooks say only "Production" — use the workbook name so they stay distinguishable.
            fl = file_label.upper()
            dept = 'Ball Valve' if 'BALL' in fl else 'Conveyor' if 'CONVEYOR' in fl else (dept or title)
        else:
            dept = dept.title()
        if title.lower() == 'quality' and 'BALL' in file_label.upper():
            flag('Skill matrix', 'low', f'{file_label} › {title}', 'Sheet is named "Quality" but holds Ball Valve production skills (dept "Production")')

        skills = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(hr, c).value
            if isinstance(v, str) and v.strip().lower() == 'expected':
                name = clean(merged_top(ws, hr - 1, c))
                group = clean(merged_top(ws, hr - 2, c)) if hr >= 3 else None
                if not name:
                    continue  # unnamed placeholder columns
                skills.append({'name': name, 'group': group, 'col': c})
        names_seen = {}
        dup = [s['name'] for s in skills]
        for i, s in enumerate(skills):
            if dup.count(s['name']) > 1:
                s['name'] = f"{s['name']} ({s['group']})" if s['group'] else s['name']

        emps = []
        for r in range(hr + 1, ws.max_row + 1):
            name = clean(ws.cell(r, nc).value)
            if not isinstance(name, str):
                continue
            if name.lower() in STOP_NAMES:
                break
            ratings = []
            for s in skills:
                e, cur = level(ws.cell(r, s['col']).value), level(ws.cell(r, s['col'] + 1).value)
                ratings.append([e, cur])
            emp = {
                'name': name,
                'sheetId': clean(ws.cell(r, cols['id']).value) if 'id' in cols else None,
                'title': clean(ws.cell(r, cols['title']).value) if 'title' in cols else None,
                'ratings': ratings,
            }
            if emp['sheetId'] is not None:
                emp['sheetId'] = str(emp['sheetId'])
            emps.append(emp)
        out.append({'dept': dept, 'source': f'{Path(path).name} › {title}', 'groups': None,
                    'skills': [{'name': s['name'], 'group': s['group']} for s in skills], 'employees': emps})
    return out


# ------------------------------------------------------------------ name matching
def norm_tokens(name):
    t = re.sub(r'[^a-z ]', ' ', name.lower())
    return [x for x in t.split() if len(x) > 1]


def stem(tok):
    # Spelling variants seen in the sheets: Amith/Amit, Ganesha/Ganesh, Shivukumar/Shivakumar, Renukaprasad/Renuka Prasad
    t = tok.replace('th', 't').replace('sh', 's').replace('v', 'w')
    t = re.sub(r'[aeiou]', '', t)
    return t


GENERIC = {'kumar', 'singh', 'rai', 'ray', 'gowda'}  # shared surnames/suffixes — never enough on their own


def candidates(name, master):
    toks = norm_tokens(name)
    if not toks:
        return []
    joined = ''.join(toks)
    first = toks[0]  # the given name the sheets use ("Velu", "Amith", "Renukaprasad")
    out = []
    for e in master:
        mt = norm_tokens(e['name'])
        pairs = {stem(''.join(mt[i:i + 2])) for i in range(len(mt))}
        if joined == ''.join(mt) \
                or (first not in GENERIC and len(first) > 2 and any(stem(first) == stem(b) for b in mt if b not in GENERIC)) \
                or stem(joined) in pairs:
            out.append(e)
    return out


def match_skills(books, master):
    for b in books:
        for e in b['employees']:
            c = candidates(e['name'], master)
            # A different initial (e.g. "MAHESH M" vs master "Mahesh V H") is a conflict, not a match.
            toks = norm_tokens(e['name'])
            initials = [t for t in re.sub(r'[^A-Za-z ]', ' ', e['name']).split() if len(t) == 1]
            if initials:
                c = [m for m in c if all(i.lower() in [x[0] for x in m['name'].lower().split()] for i in initials)
                     and any(len(x) == 1 and x.lower() in [i.lower() for i in initials] for x in m['name'].split())]
            e['code'] = None
            e['match'] = 'none'
            e['candidates'] = [m['id'] for m in c]
            if len(c) == 1:
                exact = norm_tokens(c[0]['name']) == toks
                e['code'] = c[0]['id']
                e['match'] = 'exact' if exact else 'name'
                if not exact:
                    flag('Skill matrix', 'low', f"{b['dept']} › {e['name']}",
                         f"Linked to {c[0]['id']} {c[0]['name']} by unique name match — please confirm")
            elif len(c) > 1:
                e['match'] = 'ambiguous'
                flag('Skill matrix', 'high', f"{b['dept']} › {e['name']}",
                     'Ambiguous — matches ' + ', '.join(f"{m['id']} {m['name']}" for m in c) + ' — not linked')
            else:
                hint = f" (sheet ID {e['sheetId']} is not a CAPL code)" if e['sheetId'] else ''
                flag('Skill matrix', 'high', f"{b['dept']} › {e['name']}",
                     f'Not in employee master{hint} — shown in Skills only, no code assigned')
    # same employee code twice within one dept
    for b in books:
        codes = [e['code'] for e in b['employees'] if e['code']]
        for cd in {x for x in codes if codes.count(x) > 1}:
            flag('Skill matrix', 'high', f"{b['dept']} › {cd}", 'Same employee code matched twice in one sheet')


def validate_skills(books):
    for b in books:
        if not b['employees']:
            flag('Skill matrix', 'medium', b['dept'], f"Template only — {len(b['skills'])} skills defined, no employees rated ({b['source']})")
            continue
        for e in b['employees']:
            rated = [r for r in e['ratings'] if r[1] not in (None, 'NA')]
            if not rated and all(r == [None, None] for r in e['ratings']):
                flag('Skill matrix', 'medium', f"{b['dept']} › {e['name']}", 'No ratings entered')
                continue
            blanks = [b['skills'][i]['name'] for i, r in enumerate(e['ratings']) if isinstance(r[0], int) and r[1] is None]
            if blanks:
                flag('Skill matrix', 'medium', f"{b['dept']} › {e['name']}",
                     f'{len(blanks)} skill(s) have an Expected level but no Current rating — treated as "not assessed"')
            over = [b['skills'][i]['name'] for i, r in enumerate(e['ratings'])
                    if isinstance(r[0], int) and isinstance(r[1], int) and r[1] > r[0]]
            if over:
                flag('Skill matrix', 'low', f"{b['dept']} › {e['name']}",
                     f'Current above Expected for: {", ".join(over)} — check whether columns were swapped')
            odd = [r for r in e['ratings'] for v in r if v not in (None, 'NA') and v not in (1, 2, 3, 4)]
            if odd:
                flag('Skill matrix', 'high', f"{b['dept']} › {e['name']}", f'Rating outside 1–4 scale: {odd}')


# ------------------------------------------------------------------ main
def js_module(header, name, data):
    return f'// {header}\n// GENERATED by scripts/import-data.py — do not edit by hand; re-run the import instead.\n\n' \
           f'export const {name} = {json.dumps(data, indent=1, ensure_ascii=False)};\n'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--master', required=True)
    ap.add_argument('--salary', required=True)
    ap.add_argument('--skills', nargs='+', required=True)
    ap.add_argument('--payroll-out', required=True)
    ap.add_argument('--report-out', required=True)
    a = ap.parse_args()

    master = read_master(a.master)
    payroll = read_salary(a.salary)
    books = []
    for p in sorted(a.skills):
        books += read_skill_book(p)
    match_skills(books, master)
    validate_skills(books)

    # department per employee comes only from the skill matrix they appear in
    dept_of = {}
    for b in books:
        for e in b['employees']:
            if e['code']:
                dept_of.setdefault(e['code'], []).append(b['dept'])
    for code, depts in dept_of.items():
        if len(set(depts)) > 1:
            flag('Skill matrix', 'medium', code, f'Appears in several department matrices: {sorted(set(depts))}')
    employees = [{**e, 'dept': (dept_of.get(e['id']) or [None])[0]} for e in master]

    src = {'master': Path(a.master).name, 'skills': [Path(p).name for p in sorted(a.skills)],
           'importedOn': dt.date.today().isoformat()}
    (ROOT / 'js' / 'data-employees.js').write_text(
        js_module('CAPL employee master — name + employee code are the only fields in the source file.',
                  'EMPLOYEE_MASTER', {'source': src['master'], 'importedOn': src['importedOn'], 'employees': employees}),
        encoding='utf-8')
    (ROOT / 'js' / 'data-skills.js').write_text(
        js_module('CAPL skill matrices (Expected vs Current, scale 1–4).', 'SKILL_DATA',
                  {'importedOn': src['importedOn'], 'scale': LEGEND, 'matrices': books}),
        encoding='utf-8')

    # If the register's Employee's Name column is filled in later, link rows to codes by exact name (unique only).
    payroll['map'] = {}
    by_name = {}
    for e in master:
        by_name.setdefault(' '.join(norm_tokens(e['name'])), []).append(e['id'])
    for r in payroll['rows']:
        if not r['name']:
            continue
        hits = by_name.get(' '.join(norm_tokens(r['name'])), [])
        if len(hits) == 1:
            payroll['map'][str(r['slNo'])] = hits[0]
        else:
            flag('Salary register', 'high', f"Sl {r['slNo']}",
                 f"Name {r['name']!r} " + ('matches several codes' if hits else 'not in employee master') + ' — not linked')
    payroll['employeesInMaster'] = len(master)
    payroll['importedOn'] = src['importedOn']
    Path(a.payroll_out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.payroll_out).write_text(json.dumps({'kind': 'quickair-hr-payroll', 'version': 1, **payroll}, indent=1), encoding='utf-8')
    report = {'sources': src, 'counts': {
        'employees': len(master), 'salaryRows': len(payroll['rows']),
        'skillSheets': len(books), 'skillPeople': sum(len(b['employees']) for b in books),
        'skillLinked': sum(1 for b in books for e in b['employees'] if e['code'])},
        'flags': flags}
    Path(a.report_out).write_text(json.dumps(report, indent=1, ensure_ascii=False), encoding='utf-8')
    print(json.dumps(report['counts']))
    for f in flags:
        print(f"[{f['severity']:6}] {f['area']} | {f['item']} | {f['detail']}")


if __name__ == '__main__':
    main()
